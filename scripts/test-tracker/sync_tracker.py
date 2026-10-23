#!/usr/bin/env python3
"""
Synchronise COMPONENT_TESTING_TRACKER.xlsx with what actually ran.

    python scripts/test-tracker/sync_tracker.py            # inventory + unit tests + code checks + manual results
    python scripts/test-tracker/sync_tracker.py --http     # + HTTP/API checks against the local app (or TRACKER_BASE_URL)
    python scripts/test-tracker/sync_tracker.py --lint --build
    python scripts/test-tracker/sync_tracker.py --all      # everything

Rules this script follows:
  * A status is only ever set from an execution in this run, or from a dated
    entry in manual-results.json. Nothing is marked PASS because code "looks
    right" or because a file exists.
  * IDs are stable. Rows are matched by file path (components), by
    "<test file> :: <test name>" (unit tests) or by catalogue ID, never by
    position. Nothing is deleted: vanished files/tests are marked removed/retired.
  * Every status change is appended to 'Status History'; every run to 'Run Log';
    lint/build results go to 'Build & Lint', separate from functional results.

Requires Python 3.9+ with openpyxl, and Node.js on PATH.
"""

import argparse
import datetime as dt
import json
import os
import re
import subprocess
import sys
import urllib.error
import urllib.request

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: python -m pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import catalog  # noqa: E402

ROOT = catalog.ROOT
WB_PATH = os.path.join(ROOT, "COMPONENT_TESTING_TRACKER.xlsx")
MANUAL_RESULTS = os.path.join(HERE, "manual-results.json")
EVIDENCE_ROOT = os.path.join(ROOT, "test-evidence", "runs")

BASE_URL = os.environ.get("TRACKER_BASE_URL", "http://localhost:3000").rstrip("/")
API_URL = os.environ.get("TRACKER_API_URL", "http://localhost:5000").rstrip("/")

STATUSES = ["PASS", "FAIL", "BLOCKED", "NOT RUN"]
TEST_TYPES = ["Unit", "Unit – Regression", "Code-level check", "API", "Integration (HTTP)", "UI – Browser", "E2E"]
# Where HTTP and browser runs happen: "local" for localhost targets, else "deployed".
TARGET = "local" if re.search(r"//(localhost|127\.0\.0\.1)(:|/|$)", BASE_URL) else "deployed"
EXEC_METHODS = ["Automated – node test script", "Automated – static code scan", "Automated – HTTP request (local)",
                "Automated – HTTP request (deployed)", "Manual – browser (local)", "Manual – browser (deployed)"]
SEVERITIES = ["Critical", "High", "Medium", "Low"]
DEFECT_STATUSES = ["OPEN", "FIXED", "REOPENED", "CLOSED"]
RETIRED = "[RETIRED] "
REMOVED_REQ = "N/A — file removed"

S_INV, S_TC, S_DEF, S_SUM = "Component Inventory", "Test Cases", "Defects", "Test Summary"
S_BL, S_RUN, S_HIST, S_README = "Build & Lint", "Run Log", "Status History", "README"

INV_H = ["Component ID", "Component Name", "Component Type", "Module / Feature", "File Path", "Purpose",
         "Dependencies", "Existing Tests", "Testing Required", "Last Verified", "Overall Status", "Notes"]
TC_H = ["Test Case ID", "Component ID", "Component Name", "Feature", "Test Scenario", "Preconditions", "Test Steps",
        "Expected Result", "Actual Result", "Test Type", "Execution Method", "Status", "Failure Details",
        "Evidence / Test Log", "Last Executed", "Code or Test Reference"]
DEF_H = ["Defect ID", "Related Test Case ID", "Component", "Description", "Steps to Reproduce", "Expected Behavior",
         "Actual Behavior", "Severity", "Status", "Related File", "Fix Reference", "Retest Result", "Last Updated"]
BL_H = ["Run ID", "Check", "Command", "Executed At", "Result", "Exit Code", "Errors", "Warnings", "Summary",
        "Evidence / Log", "Git Commit"]
RUN_H = ["Run ID", "Started", "Finished", "Git Branch", "Git Commit", "Working Tree", "Scope", "Unit PASS",
         "Unit FAIL", "Code checks PASS", "Code checks FAIL", "HTTP PASS", "HTTP FAIL", "HTTP BLOCKED",
         "Manual results applied", "Notes"]
HIST_H = ["Timestamp", "Run ID", "Test Case ID", "Previous Status", "New Status", "Source", "Note"]


# ==========================================================================
# small utilities
# ==========================================================================

def now():
    return dt.datetime.now().replace(microsecond=0)


def rel(path):
    return os.path.relpath(path, ROOT).replace(os.sep, "/")


def read_text(rel_path):
    with open(os.path.join(ROOT, rel_path), encoding="utf-8", errors="replace") as fh:
        return fh.read()


def git(*args):
    try:
        return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, timeout=30).stdout.strip()
    except Exception:
        return ""


def as_dt(value):
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return dt.datetime.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def clip(text, n=1500):
    text = "" if text is None else str(text)
    return text if len(text) <= n else text[: n - 15] + "\n…[truncated]"


# ==========================================================================
# workbook I/O
# ==========================================================================

def read_rows(wb, sheet):
    if wb is None or sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not values or values[0] in (None, ""):
            continue
        rows.append({h: v for h, v in zip(headers, values) if h})
    return rows


# ==========================================================================
# 1. component inventory
# ==========================================================================

INV_DIRS = ["src"]
INV_EXTRA = ["scripts/ui-smoke.mjs"]
CODE_EXT = (".js", ".jsx", ".ts", ".tsx", ".mjs")
TYPE_ORDER = ["Middleware", "Page", "Layout", "Route boundary", "Page module", "API Route", "Service",
              "Hook (React Query)", "Hook", "Context", "Provider", "Feature module", "Lib / Utility", "Component",
              "Constants", "Types", "Test tooling"]
TYPE_REQ = {
    "Middleware": "Integration (HTTP)", "Page": "UI – Browser / E2E", "Layout": "UI smoke",
    "Route boundary": "UI smoke", "Page module": "Component / UI", "API Route": "Integration (route handler)",
    "Service": "API contract + integration", "Hook (React Query)": "Unit (mocked service)", "Hook": "Unit",
    "Context": "Unit / integration", "Provider": "Smoke", "Feature module": "Unit + UI", "Lib / Utility": "Unit",
    "Component": "Component / UI", "Constants": "None (static data)", "Types": "None (types only)",
    "Test tooling": "N/A (test tooling)",
}
HIGH_RX = re.compile(r"auth|login|register|otp|password|middleware|axios|payment|razorpay|pricing|store|enrol"
                     r"|courseMapper|progressIndex|reorder|/learn|composer|attempt|quiz|publish", re.I)
SPECIAL_TITLE = {"qa": "Q&A", "ai": "AI", "tts": "TTS", "api": "API", "otp": "OTP"}


def title(seg):
    seg = re.sub(r"\.(jsx?|tsx?|mjs)$", "", seg)
    if seg in SPECIAL_TITLE:
        return SPECIAL_TITLE[seg]
    if seg.startswith("["):
        return seg
    return " ".join(w[:1].upper() + w[1:] for w in re.split(r"[-_]", seg) if w)


def classify(path):
    parts = path.split("/")
    name = parts[-1]
    stem = re.sub(r"\.(jsx?|tsx?|mjs)$", "", name)
    if path.startswith("scripts/"):
        return "Test tooling", "QA tooling", stem
    if path == "src/middleware.js":
        return "Middleware", "Auth / Route guard", "middleware (route guard)"
    area = parts[1]
    if area == "app":
        segs = parts[2:]
        route = "/" + "/".join(segs[:-1])
        if segs[0] == "api":
            return "API Route", f"Next API routes / {title(segs[1])}", route
        kind = {"page": "Page", "layout": "Layout", "loading": "Route boundary", "error": "Route boundary",
                "not-found": "Route boundary"}.get(stem, "Page module")
        first = segs[0] if len(segs) > 1 else ""
        if first in ("admin", "instructor", "student"):
            sub = title(segs[1]) if len(segs) > 2 else "Home"
            feature = f"{title(first)} / {sub}"
        elif first == "courses":
            feature = "Public / Courses"
        elif first in ("login", "register", "verify-otp", "forgot-password", "reset-password"):
            feature = "Auth"
        elif first in ("zz-slider-check", "qa-course-preview"):
            feature = "Dev / Temporary harness"
        elif first == "dashboard":
            feature = "Public / Dashboard"
        else:
            feature = "Public / Landing"
        label = route if kind == "Page" else f"{route} ({stem})"
        return kind, feature, label
    if area == "components":
        group = parts[2]
        sub = parts[3] if len(parts) > 4 else "General"
        if group in ("admin", "instructor", "student"):
            feature = f"{title(group)} / {title(sub)}"
        elif group in ("ui", "common", "shared", "forms", "menus", "modules"):
            feature = "Shared UI"
        elif group == "auth":
            feature = "Auth"
        elif group == "home":
            feature = "Public / Landing"
        elif group in ("courses", "course-details"):
            feature = "Public / Courses"
        else:
            feature = title(group)
        return "Component", feature, stem if stem != "index" else parts[-2]
    if area == "hooks":
        if parts[2] == "queries":
            role = title(parts[3]) if len(parts) > 4 else "Public"
            return "Hook (React Query)", f"{role} / Data hooks", stem
        return "Hook", "Shared hooks", stem
    if area == "services":
        return "Service", "API services", stem
    if area == "lib":
        return "Lib / Utility", ("Shared lib / External files" if parts[2] == "external" else "Shared lib"), stem
    if area == "utils":
        return "Lib / Utility", "Shared lib", stem
    if area == "context":
        return "Context", "App state", stem
    if area == "providers":
        return "Provider", "App shell", stem
    if area == "features":
        return "Feature module", f"Feature / {title(parts[2])}", stem
    if area == "constants":
        return "Constants", "Shared lib", stem
    if area == "types":
        return "Types", "Shared lib", stem
    return "Component", title(area), stem


def list_source_files():
    out = []
    for d in INV_DIRS:
        for dirpath, dirnames, filenames in os.walk(os.path.join(ROOT, d)):
            dirnames[:] = sorted(x for x in dirnames if x != "__tests__")
            for f in sorted(filenames):
                if f.endswith(CODE_EXT) and not f.endswith(".d.ts"):
                    out.append(rel(os.path.join(dirpath, f)))
    out += [p for p in INV_EXTRA if os.path.exists(os.path.join(ROOT, p))]
    return out


_IMPORT_RX = re.compile(r"""(?:import\s[^;]*?from\s*|import\s*\(\s*|require\(\s*)[`"']([^`"']+)[`"']""", re.S)


def imports_of(src):
    return _IMPORT_RX.findall(src)


def resolve_import(from_rel, spec):
    spec = spec.split("?")[0]
    if spec.startswith("@/"):
        base = os.path.join(ROOT, "src", spec[2:])
    elif spec.startswith("."):
        base = os.path.join(ROOT, os.path.dirname(from_rel), spec)
    else:
        return None
    base = os.path.normpath(base)
    for cand in [base] + [base + e for e in CODE_EXT] + [os.path.join(base, "index" + e) for e in CODE_EXT]:
        if os.path.isfile(cand):
            return rel(cand)
    return None


def file_header(src):
    """Text before the first line of real code (directives, imports and comments only)."""
    out, in_import, in_block = [], False, False
    for line in src.splitlines():
        s = line.strip()
        if in_block:
            out.append(line)
            in_block = "*/" not in s
            continue
        if in_import:
            in_import = not re.search(r"""from\s*["'][^"']+["']|^["'][^"']+["'];?$""", s)
            continue
        if not s or re.match(r"""^["']use (client|server)["'];?$""", s):
            continue
        if s.startswith("/*"):
            out.append(line)
            in_block = "*/" not in s
            continue
        if s.startswith("//"):
            out.append(line)
            continue
        if s.startswith("import "):
            in_import = not re.search(r"""from\s*["'][^"']+["']|^import\s+["'][^"']+["']""", s)
            continue
        break
    return "\n".join(out)


def doc_sentence(src):
    """First explanatory comment in the file header (never a comment on an inner helper)."""
    head = file_header(src)
    blocks = re.findall(r"/\*\*?(.*?)\*/|((?:^[ \t]*//[^\n]*\n?)+)", head, re.S | re.M)
    for block, line_block in blocks:
        text = block or line_block
        text = re.sub(r"^[ \t]*(\*|//)+ ?", "", text, flags=re.M)
        text = "\n".join(l for l in text.splitlines() if not l.strip().startswith("@"))
        text = " ".join(re.sub(r"[=\-*#~_]{3,}", " ", text).split())
        if len(text) < 25 or re.match(r"(eslint|@ts-|prettier|biome|#region|TODO)", text, re.I):
            continue
        m = re.match(r"(.{25,220}?[.!?])(\s|$)", text)
        return (m.group(1) if m else text[:220]).strip()
    return ""


def purpose_for(path, kind, label, src):
    if path in catalog.PURPOSE_OVERRIDES:
        return catalog.PURPOSE_OVERRIDES[path]
    doc = doc_sentence(src)
    if kind == "Page":
        return f"Route {label}." + (f" {doc}" if doc else "")
    if kind == "Service":
        eps = []
        for m, p in re.findall(r"api\.(get|post|put|patch|delete)\(\s*[`\"']([^`\"']+)", src):
            ep = f"{m.upper()} {re.sub(r'\$\{[^}]+\}', ':param', p)}"
            if ep not in eps:
                eps.append(ep)
        if eps:
            more = f" (+{len(eps) - 6} more)" if len(eps) > 6 else ""
            return "API client: " + ", ".join(eps[:6]) + more
    if kind == "Hook (React Query)":
        kinds = sorted(set(re.findall(r"\b(useQuery|useMutation|useInfiniteQuery|useQueries)\b", src)))
        fns = []
        for names, spec in re.findall(r"import\s*\{([^}]*)\}\s*from\s*[\"'](@/services/[^\"']+)[\"']", src):
            fns += [n.strip().split(" as ")[0] for n in names.split(",") if n.strip()]
        if kinds or fns:
            return f"React Query {'/'.join(kinds) or 'hook'}" + (f" over {', '.join(fns[:5])}" if fns else "") + (f". {doc}" if doc else "")
    if doc:
        return doc
    return {"Component": "UI component", "Hook": "Shared React hook", "Lib / Utility": "Utility module",
            "Context": "React context provider", "Provider": "App-level provider", "Feature module": "Feature module",
            "Constants": "Constant definitions", "Types": "Type definitions", "Layout": "Route layout",
            "API Route": "Next.js route handler", "Page module": "Route-local module"}.get(kind, kind) + f" ({label}) — purpose inferred from path."


def deps_for(path, src):
    names = []
    for spec in imports_of(src):
        target = resolve_import(path, spec)
        if target and target != path:
            n = re.sub(r"\.(jsx?|tsx?|mjs)$", "", target.split("/")[-1])
            if n == "index":
                n = target.split("/")[-2]
            if n not in names:
                names.append(n)
    if not names:
        return "—"
    return ", ".join(names[:10]) + (f" (+{len(names) - 10} more)" if len(names) > 10 else "")


def test_files():
    out = []
    for dirpath, _, filenames in os.walk(os.path.join(ROOT, "src")):
        for f in sorted(filenames):
            if f.endswith(".test.js"):
                out.append(rel(os.path.join(dirpath, f)))
    return sorted(out)


def test_import_map(tests):
    m = {}
    for t in tests:
        for spec in imports_of(read_text(t)):
            target = resolve_import(t, spec)
            if target and "__tests__" not in target:
                m.setdefault(target, []).append(t)
    return m


def build_inventory(existing, tests, run_date):
    tmap = test_import_map(tests)
    by_path = {r["File Path"]: r for r in existing}
    max_id = max([int(re.sub(r"\D", "", r["Component ID"]) or 0) for r in existing] + [0])
    files = list_source_files()
    present = set(files)
    rows = []
    fresh = []
    for p in files:
        kind, feature, label = classify(p)
        src = read_text(p)
        req = TYPE_REQ.get(kind, "Component / UI")
        if not req.startswith(("None", "N/A")):
            prio = "High" if HIGH_RX.search(p) else ("Low" if feature == "Shared UI" or kind in ("Constants", "Types") else "Medium")
            req = f"{req} — {prio}"
        row = by_path.get(p)
        tests_for = tmap.get(p, [])
        computed = {
            "Component Name": label, "Component Type": kind, "Module / Feature": feature, "File Path": p,
            "Dependencies": deps_for(p, src),
            "Existing Tests": ", ".join(t.split("/")[-1] for t in tests_for) if tests_for else "None",
        }
        if row:
            row.update(computed)
            # A Purpose starting with "=" came from a banner comment that an
            # earlier version wrote as a formula; regenerate it.
            if not row.get("Purpose") or str(row["Purpose"]).startswith("="):
                row["Purpose"] = purpose_for(p, kind, label, src)
            if not row.get("Testing Required") or row.get("Testing Required") == REMOVED_REQ:
                row["Testing Required"] = req
            if str(row.get("Notes") or "").startswith("REMOVED"):
                row["Notes"] = f"RESTORED {run_date}: file exists again."
        else:
            row = dict(computed, Purpose=purpose_for(p, kind, label, src), **{"Testing Required": req, "Notes": ""})
            fresh.append((TYPE_ORDER.index(kind) if kind in TYPE_ORDER else 99, p, row))
        rows.append(row)
    for _, _, row in sorted(fresh, key=lambda x: (x[0], x[1])):
        max_id += 1
        row["Component ID"] = f"CMP-{max_id:04d}"
    for r in existing:
        if r["File Path"] not in present:
            if not str(r.get("Notes") or "").startswith("REMOVED"):
                r["Notes"] = f"REMOVED {run_date}: file no longer exists. " + str(r.get("Notes") or "")
            r["Testing Required"] = REMOVED_REQ
            r["Existing Tests"] = r.get("Existing Tests") or "None"
            rows.append(r)
    rows.sort(key=lambda r: r["Component ID"])
    return rows


# ==========================================================================
# 2. executions
# ==========================================================================

UNIT_OVERRIDES = [
    ("textToSpeech", re.compile(r"^(question speech|result summary speech|answer review speech)"), "src/lib/quizSpeech.js"),
    ("videoPlayback", re.compile(r"^(watch URL|mobile watch|short youtu|embed URL|privacy|Shorts URL|live URL|invalid:|host:)"),
     "src/lib/youtube.js"),
]


def unit_component(test_rel, name, tmap_rev):
    stem = test_rel.split("/")[-1].replace(".test.js", "")
    for s, rx, target in UNIT_OVERRIDES:
        if s == stem and rx.search(name):
            return target
    test_dir = os.path.dirname(os.path.dirname(test_rel))
    for cand in (f"{test_dir}/{stem}.js", f"src/lib/{stem}.js"):
        if os.path.exists(os.path.join(ROOT, cand)):
            return cand
    return tmap_rev.get(test_rel, [None])[0]


def run_unit_tests(tests, run_dir):
    os.makedirs(os.path.join(run_dir, "unit"), exist_ok=True)
    results = []  # dicts: file, name, status, line, details, log, exit
    files = {}
    for t in tests:
        stem = t.split("/")[-1].replace(".test.js", "")
        started = now()
        try:
            proc = subprocess.run(["node", t], cwd=ROOT, capture_output=True, text=True, timeout=600,
                                  encoding="utf-8", errors="replace")
            out, code = (proc.stdout or "") + (proc.stderr or ""), proc.returncode
        except Exception as exc:
            out, code = f"could not execute: {exc}", None
        log_path = os.path.join(run_dir, "unit", f"{stem}.log")
        with open(log_path, "w", encoding="utf-8") as fh:
            fh.write(f"$ node {t}\n# started {started.isoformat()}  exit={code}\n\n{out}")
        lines = out.splitlines()
        found = []
        for i, line in enumerate(lines):
            m = re.match(r"^\[(PASS|FAIL)\]\s+(.*)$", line.strip())
            if not m:
                continue
            details = []
            for nxt in lines[i + 1: i + 7]:
                if re.match(r"^\[(PASS|FAIL)\]|^RESULTS", nxt.strip()):
                    break
                if nxt.strip():
                    details.append(nxt.strip())
            found.append(dict(file=t, name=m.group(2).strip(), status=m.group(1), line=i + 4,
                              details="\n".join(details) if m.group(1) == "FAIL" else "", log=rel(log_path)))
        results += found
        files[t] = dict(exit=code, log=rel(log_path), reported={f["name"] for f in found}, tail="\n".join(lines[-8:]))
    return results, files


def run_code_checks(run_dir):
    out = []
    for c in catalog.CODE_CHECKS:
        if c.get("retired"):
            continue
        try:
            status, actual, evidence = c["fn"]()
        except Exception as exc:
            status, actual, evidence = "BLOCKED", "", f"check raised {type(exc).__name__}: {exc}"
        out.append(dict(c, status=status, actual=actual, evidence=evidence))
    with open(os.path.join(run_dir, "code-checks.json"), "w", encoding="utf-8") as fh:
        json.dump([{k: v for k, v in r.items() if k != "fn"} for r in out], fh, indent=2, ensure_ascii=False)
    return out


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


def run_http_checks(run_dir):
    opener = urllib.request.build_opener(_NoRedirect)
    out = []
    for c in catalog.HTTP_CHECKS:
        if c.get("retired"):
            continue
        base = BASE_URL if c["target"] == "web" else API_URL
        url = base + c["path"]
        data = json.dumps(c["json"]).encode() if "json" in c else None
        headers = {"User-Agent": "orange-lms-test-tracker/1.0", **c.get("headers", {})}
        if data is not None:
            headers["Content-Type"] = "application/json"
        req = urllib.request.Request(url, data=data, method=c["method"], headers=headers)
        status, loc, body, err = None, "", "", ""
        try:
            with opener.open(req, timeout=90) as resp:
                status, loc, body = resp.status, resp.headers.get("Location", ""), resp.read(200000).decode("utf-8", "replace")
        except urllib.error.HTTPError as e:
            status, loc = e.code, e.headers.get("Location", "")
            body = e.read(200000).decode("utf-8", "replace")
        except Exception as exc:
            err = f"{type(exc).__name__}: {exc}"
        problems = []
        extra = ""
        if err:
            res = "BLOCKED"
        else:
            if status != c["expect_status"]:
                problems.append(f"status {status} != expected {c['expect_status']}")
            if "expect_location" in c:
                norm = loc.replace(base, "") or "/"
                if norm != c["expect_location"]:
                    problems.append(f"Location '{loc}' != '{c['expect_location']}'")
            if "expect_body" in c and c["expect_body"].lower() not in body.lower():
                problems.append(f"body does not contain '{c['expect_body']}'")
            if c.get("expect_json") == "every_item_has_store_key":
                try:
                    d = json.loads(body)
                    d = d.get("data", d)
                    items = d if isinstance(d, list) else (d.get("courses") or d.get("items") or d.get("data") or [])
                    with_key = [x for x in items if "store" in x]
                    nulls = sum(1 for x in with_key if x.get("store") is None)
                    extra = f"{len(items)} item(s); {len(with_key)} carry 'store'; {nulls} with store=null"
                    if not items or len(with_key) != len(items):
                        problems.append("not every item has a 'store' key")
                except Exception as exc:
                    problems.append(f"invalid JSON: {exc}")
            res = "FAIL" if problems else "PASS"
        req_desc = f"{c['method']} {url}" + (f"  headers={c.get('headers')}" if c.get("headers") else "") + \
                   (f"  body={json.dumps(c['json'])}" if "json" in c else "")
        resp_desc = err or (f"HTTP {status}" + (f"; Location: {loc}" if loc else "") +
                            (f"; {extra}" if extra else "") + f"; body[0:160]: {body[:160]!r}")
        out.append(dict(c, status=res, request=req_desc, response=resp_desc,
                        actual=err and "" or resp_desc, failure="; ".join(problems) if problems else (err and f"Endpoint unreachable: {err}" or "")))
    with open(os.path.join(run_dir, "http-checks.json"), "w", encoding="utf-8") as fh:
        json.dump([{k: v for k, v in r.items()} for r in out], fh, indent=2, ensure_ascii=False)
    return out


def _npm(args, log_path):
    cmd = "npm " + args
    started = now()
    proc = subprocess.run(cmd, cwd=ROOT, shell=True, capture_output=True, text=True, encoding="utf-8", errors="replace")
    with open(log_path, "w", encoding="utf-8") as fh:
        fh.write(f"$ {cmd}\n# started {started.isoformat()} exit={proc.returncode}\n\n{proc.stdout}\n{proc.stderr}")
    return proc, started, cmd


def run_lint(run_dir, run_id, commit):
    report = os.path.join(run_dir, "eslint.json")
    proc, started, cmd = _npm(f'run lint -- -f json -o "{report}"', os.path.join(run_dir, "lint.log"))
    errors = warnings = None
    summary = ""
    try:
        data = json.load(open(report, encoding="utf-8"))
        errors = sum(f["errorCount"] for f in data)
        warnings = sum(f["warningCount"] for f in data)
        err_files = [rel(f["filePath"]) for f in data if f["errorCount"]]
        src_err = sum(f["errorCount"] for f in data if rel(f["filePath"]).startswith("src/"))
        summary = f"{errors} error(s) in {len(err_files)} file(s) ({src_err} under src/); {warnings} warning(s)."
        if err_files:
            summary += " Error files: " + ", ".join(err_files[:12]) + (" …" if len(err_files) > 12 else "")
    except Exception as exc:
        summary = f"No ESLint JSON report ({exc}); see log."
    return dict(zip(BL_H, [run_id, "Lint", cmd.replace(report, "<run>/eslint.json"), started,
                           "PASS" if proc.returncode == 0 else "FAIL", proc.returncode, errors, warnings, summary,
                           rel(os.path.join(run_dir, "lint.log")) + "; " + rel(report), commit]))


def run_build(run_dir, run_id, commit):
    log = os.path.join(run_dir, "build.log")
    proc, started, cmd = _npm("run build", log)
    text = proc.stdout + proc.stderr
    pages = re.findall(r"Generating static pages[^(]*\((\d+)/(\d+)\)", text)
    summary = ("Build succeeded" if proc.returncode == 0 else "Build FAILED") + \
              (f"; static pages {pages[-1][0]}/{pages[-1][1]}" if pages else "")
    if proc.returncode != 0:
        summary += "\n" + "\n".join(text.strip().splitlines()[-12:])
    return dict(zip(BL_H, [run_id, "Build", cmd, started, "PASS" if proc.returncode == 0 else "FAIL",
                           proc.returncode, None, None, clip(summary, 900), rel(log), commit]))


# ==========================================================================
# 3. merge results into test cases / defects
# ==========================================================================

def merge(existing_tc, inventory, unit_results, unit_files, code_results, http_results, manual_log,
          run_id, run_ts, hist, tests, existing_def, lint_row):
    inv_by_path = {r["File Path"]: r for r in inventory}
    tcs = {r["Test Case ID"]: r for r in existing_tc}
    by_ref = {r.get("Code or Test Reference"): r for r in existing_tc if str(r["Test Case ID"]).startswith("TC-UT-")}
    tmap_rev = {}
    for t in tests:
        for spec in imports_of(read_text(t)):
            target = resolve_import(t, spec)
            if target and "__tests__" not in target:
                tmap_rev.setdefault(t, []).append(target)

    def link(row, comp_path):
        comp = inv_by_path.get(comp_path)
        row["Component ID"] = comp["Component ID"] if comp else "UNMAPPED"
        row["Component Name"] = comp["Component Name"] if comp else comp_path
        row["Feature"] = comp["Module / Feature"] if comp else "—"

    def set_status(row, status, source, note=""):
        prev = row.get("Status") or "—"
        row["Status"] = status
        if prev != status:
            hist.append(dict(zip(HIST_H, [run_ts, run_id, row["Test Case ID"], prev, status, source, clip(note, 300)])))

    def unretire(row):
        if str(row.get("Test Scenario") or "").startswith(RETIRED):
            row["Test Scenario"] = row["Test Scenario"][len(RETIRED):]

    def retire(row, why):
        if not str(row.get("Test Scenario") or "").startswith(RETIRED):
            row["Test Scenario"] = RETIRED + str(row.get("Test Scenario") or "")
            row["Failure Details"] = why
            set_status(row, "NOT RUN", "sync", why)

    # ---- unit tests ----
    counters = {}
    for r in existing_tc:
        m = re.match(r"TC-UT-([A-Z0-9]+)-(\d+)$", str(r["Test Case ID"]))
        if m:
            counters[m.group(1)] = max(counters.get(m.group(1), 0), int(m.group(2)))
    for u in unit_results:
        ref = f"{u['file']} :: {u['name']}"
        row = by_ref.get(ref)
        if row is None:
            key = re.sub(r"[^A-Z0-9]", "", u["file"].split("/")[-1].replace(".test.js", "").upper())
            counters[key] = counters.get(key, 0) + 1
            row = {"Test Case ID": f"TC-UT-{key}-{counters[key]:03d}", "Code or Test Reference": ref}
            tcs[row["Test Case ID"]] = row
            by_ref[ref] = row
        unretire(row)
        link(row, unit_component(u["file"], u["name"], tmap_rev))
        is_reg = u["name"].upper().startswith("REGRESSION")
        row.update({
            "Test Scenario": u["name"],
            "Preconditions": "Node.js ≥ 22 (ES modules via syntax detection). No network, backend or browser required.",
            "Test Steps": f"node {u['file']}",
            "Expected Result": f"Assertion holds: \"{u['name']}\" (script prints [PASS]).",
            "Actual Result": f"[{u['status']}] reported by test script",
            "Test Type": "Unit – Regression" if is_reg else "Unit",
            "Execution Method": "Automated – node test script",
            "Failure Details": clip(u["details"], 800),
            "Evidence / Test Log": f"{u['log']} (line {u['line']}); run {run_id}",
            "Last Executed": run_ts,
        })
        set_status(row, u["status"], "unit test run", u["log"])
    # vanished or unreported unit tests
    for ref, row in list(by_ref.items()):
        file = ref.split(" :: ")[0]
        name = ref.split(" :: ", 1)[1] if " :: " in ref else ""
        info = unit_files.get(file)
        if info is None:
            if not os.path.exists(os.path.join(ROOT, file)):
                retire(row, f"RETIRED {run_ts:%Y-%m-%d}: test file {file} no longer exists.")
            continue
        if name in info["reported"]:
            continue
        if info["exit"] == 0:
            retire(row, f"RETIRED {run_ts:%Y-%m-%d}: '{name}' is no longer reported by {file} (renamed or removed).")
        else:
            row.update({"Actual Result": f"Not reported — test file exited with code {info['exit']} before this assertion.",
                        "Failure Details": clip(info["tail"], 800), "Evidence / Test Log": f"{info['log']}; run {run_id}",
                        "Last Executed": run_ts})
            set_status(row, "FAIL", "unit test run", f"test file crashed (exit {info['exit']})")

    # ---- code-level checks ----
    for c in code_results:
        row = tcs.setdefault(c["id"], {"Test Case ID": c["id"]})
        unretire(row)
        link(row, c["component"])
        row.update({
            "Test Scenario": c["scenario"],
            "Preconditions": "Working tree of the frontend repository (static scan; app not executed).",
            "Test Steps": "python scripts/test-tracker/sync_tracker.py (runs catalog check " + c["fn"].__name__ + ")",
            "Expected Result": c["expected"], "Actual Result": clip(c["actual"], 800),
            "Test Type": "Code-level check", "Execution Method": "Automated – static code scan",
            "Failure Details": clip(c["evidence"], 1200) if c["status"] != "PASS" else "",
            "Evidence / Test Log": clip(c["evidence"], 1200) + f"\nrun {run_id}: test-evidence/runs/{run_id}/code-checks.json",
            "Last Executed": run_ts, "Code or Test Reference": c["ref"],
        })
        set_status(row, c["status"], "code-level check", c["actual"])

    # ---- HTTP / API checks ----
    for h in http_results:
        row = tcs.setdefault(h["id"], {"Test Case ID": h["id"]})
        unretire(row)
        link(row, h["component"])
        target = BASE_URL if h["target"] == "web" else API_URL
        exp = f"HTTP {h['expect_status']}" + (f", Location {h['expect_location']}" if "expect_location" in h else "") + \
              (f", body contains '{h['expect_body']}'" if "expect_body" in h else "") + \
              ("; every item has a 'store' key" if h.get("expect_json") else "")
        row.update({
            "Test Scenario": h["scenario"],
            "Preconditions": f"{target} reachable. Read-only / rejected-before-side-effect request.",
            "Test Steps": h["request"], "Expected Result": exp, "Actual Result": clip(h["response"], 600),
            "Test Type": "Integration (HTTP)" if h["target"] == "web" else "API",
            "Execution Method": f"Automated – HTTP request ({TARGET})",
            "Failure Details": h["failure"], "Evidence / Test Log": f"Request: {h['request']}\nResponse: {clip(h['response'], 400)}\nrun {run_id}: test-evidence/runs/{run_id}/http-checks.json",
            "Last Executed": run_ts, "Code or Test Reference": h["component"],
        })
        set_status(row, h["status"], "HTTP check", h["failure"] or h["response"][:120])
    for c in catalog.HTTP_CHECKS:  # defined but not executed this run
        if c["id"] not in tcs:
            row = {"Test Case ID": c["id"]}
            link(row, c["component"])
            target = BASE_URL if c["target"] == "web" else API_URL
            row.update({"Test Scenario": c["scenario"], "Preconditions": f"{target} reachable.",
                        "Test Steps": f"{c['method']} {target}{c['path']}", "Expected Result": f"HTTP {c['expect_status']}",
                        "Test Type": "Integration (HTTP)" if c["target"] == "web" else "API",
                        "Execution Method": f"Automated – HTTP request ({TARGET})", "Status": "NOT RUN",
                        "Code or Test Reference": c["component"]})
            tcs[c["id"]] = row

    # ---- manual / browser cases ----
    for m in catalog.MANUAL_CASES:
        row = tcs.get(m["id"])
        if row is None:
            row = {"Test Case ID": m["id"], "Status": "NOT RUN", "Actual Result": "", "Failure Details": "",
                   "Evidence / Test Log": "", "Last Executed": None}
            tcs[m["id"]] = row
        link(row, m["component"])
        row.update({"Test Scenario": (RETIRED if m.get("retired") else "") + m["scenario"], "Preconditions": m["pre"],
                    "Test Steps": m["steps"], "Expected Result": m["expected"], "Test Type": m["type"],
                    "Execution Method": m.get("method", f"Manual – browser ({TARGET})"),
                    "Code or Test Reference": m["ref"]})
    applied = 0
    latest = {}
    for e in manual_log:
        ts = as_dt(e.get("executed_at"))
        if e.get("id") and ts and (e["id"] not in latest or ts > latest[e["id"]][0]):
            latest[e["id"]] = (ts, e)
    for tc_id, (ts, e) in latest.items():
        row = tcs.get(tc_id)
        if row is None:
            print(f"  ! manual-results.json references unknown test case {tc_id}")
            continue
        last = as_dt(row.get("Last Executed"))
        if last and last >= ts:
            continue
        status = e.get("status", "NOT RUN")
        if status not in STATUSES:
            print(f"  ! {tc_id}: invalid status {status!r} ignored")
            continue
        if status in ("PASS", "FAIL") and not (e.get("evidence") or e.get("actual")):
            print(f"  ! {tc_id}: {status} without evidence ignored")
            continue
        row.update({"Actual Result": e.get("actual", ""), "Failure Details": e.get("failure_details", ""),
                    "Evidence / Test Log": "; ".join(x for x in [e.get("evidence", ""), f"by {e.get('executed_by', '?')}",
                                                                   f"env: {e.get('environment', '?')}"] if x),
                    "Last Executed": ts})
        set_status(row, status, "manual-results.json", e.get("failure_details") or e.get("actual", "")[:120])
        applied += 1
    every = catalog.CODE_CHECKS + catalog.HTTP_CHECKS + catalog.MANUAL_CASES
    live_ids = {c["id"] for c in every if not c.get("retired")}
    for tc_id, row in tcs.items():
        if not tc_id.startswith("TC-UT-") and tc_id not in live_ids:
            retire(row, f"RETIRED {run_ts:%Y-%m-%d}: no longer defined (or marked retired) in scripts/test-tracker/catalog.py.")

    # ---- defects ----
    defects = {d["Defect ID"]: d for d in existing_def}
    for s in catalog.SEED_DEFECTS:
        if s["id"] not in defects:
            defects[s["id"]] = dict(zip(DEF_H, [s["id"], s["tc"], s["component"], s["desc"], s["steps"], s["expected"],
                                                s["actual"], s["severity"], "OPEN", s["file"], "", "NOT RUN", run_ts]))
    executed_now = {r["Test Case ID"]: r["Status"] for r in tcs.values()
                    if r.get("Last Executed") == run_ts and str(r["Test Case ID"]).startswith(("TC-UT-", "TC-CL-", "TC-API-"))}
    if lint_row:
        executed_now["Build & Lint: Lint"] = lint_row["Result"]
    for d in defects.values():
        linked = re.findall(r"TC-[A-Z]+-[A-Z0-9]+(?:-\d+)?|Build & Lint: Lint", str(d.get("Related Test Case ID") or ""))
        auto = [l for l in linked if l in executed_now]
        if not auto:
            continue
        results = [executed_now[l] for l in auto]
        st = d.get("Status")
        if all(r == "PASS" for r in results):
            if st in ("OPEN", "REOPENED"):
                d["Status"] = "FIXED"
                d["Fix Reference"] = (str(d.get("Fix Reference") or "") +
                                      f"\nAuto-verified: {', '.join(auto)} PASS in run {run_id} @ {git('rev-parse', '--short', 'HEAD')}").strip()
            if d["Status"] in ("FIXED",) and d.get("Retest Result") != "PASS":
                d["Retest Result"] = "PASS"
                d["Last Updated"] = run_ts
        elif any(r == "FAIL" for r in results):
            if st in ("FIXED", "CLOSED"):
                d["Status"] = "REOPENED"
                d["Retest Result"] = "FAIL"
                d["Last Updated"] = run_ts
    # auto-log new failures that no defect covers
    covered = set()
    for d in defects.values():
        covered |= set(re.findall(r"TC-[A-Z]+-[A-Z0-9]+(?:-\d+)?", str(d.get("Related Test Case ID") or "")))
    next_def = max([int(re.sub(r"\D", "", k) or 0) for k in defects] + [0])
    for tc_id, st in executed_now.items():
        if st == "FAIL" and tc_id.startswith("TC-") and tc_id not in covered:
            r = tcs[tc_id]
            next_def += 1
            did = f"DEF-{next_def:03d}"
            defects[did] = dict(zip(DEF_H, [did, tc_id, r.get("Component Name"), f"Auto-logged by sync: '{r.get('Test Scenario')}' failed.",
                                            r.get("Test Steps"), r.get("Expected Result"), clip(r.get("Failure Details") or r.get("Actual Result"), 600),
                                            "Medium", "OPEN", r.get("Code or Test Reference"), "", "NOT RUN", run_ts]))
            print(f"  + auto-logged {did} for failing {tc_id}")

    def tc_sort(r):
        tid = r["Test Case ID"]
        order = {"TC-UT": 0, "TC-CL": 1, "TC-API": 2, "TC-UI": 3}
        return (order.get(tid[:5], 9), tid)
    return sorted(tcs.values(), key=tc_sort), sorted(defects.values(), key=lambda d: d["Defect ID"]), applied


# ==========================================================================
# 4. write workbook
# ==========================================================================

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=ARIAL, size=10)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
FILLS = {
    "PASS": ("C6EFCE", "006100"), "FAIL": ("FFC7CE", "9C0006"), "BLOCKED": ("FFEB9C", "9C5700"),
    "NOT RUN": ("EDEDED", "404040"), "PARTIAL": ("DDEBF7", "1F4E78"), "NO TEST CASES": ("F2F2F2", "7F7F7F"),
    "OPEN": ("FFC7CE", "9C0006"), "REOPENED": ("F8CBAD", "833C0B"), "FIXED": ("DDEBF7", "1F4E78"),
    "CLOSED": ("C6EFCE", "006100"), "Critical": ("C00000", "FFFFFF"), "High": ("FFC7CE", "9C0006"),
    "Medium": ("FFEB9C", "9C5700"), "Low": ("EDEDED", "404040"),
}


def status_cf(ws, rng, values):
    for v in values:
        bg, fg = FILLS[v]
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'],
                                                      fill=PatternFill("solid", fgColor=bg), font=Font(name=ARIAL, color=fg, bold=True)))


def dropdown(ws, rng, values, prompt):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid value", error="Choose a value from the list.",
                        promptTitle="Allowed values", prompt=prompt)
    dv.add(rng)
    ws.add_data_validation(dv)


def write_table(ws, headers, rows, widths, wrap=(), dates=(), freeze="B2"):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append([r.get(h) if isinstance(r, dict) else r[i] for i, h in enumerate(headers)])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            h = headers[cell.column - 1]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # data, never a formula (e.g. "==== banner ====" comments)
            cell.font = BODY_FONT
            cell.alignment = TOP_WRAP if h in wrap else TOP
            cell.border = THIN
            if h in dates:
                cell.number_format = "yyyy-mm-dd hh:mm"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"


def write_workbook(inventory, tcs, defects, bl_rows, run_rows, hist_rows, meta):
    wb = Workbook()
    n_tc = max(2000, 2 * len(tcs) + 2)
    n_inv = max(2000, 2 * len(inventory) + 2)
    n_def = max(500, 2 * len(defects) + 2)
    TC = f"'{S_TC}'"
    rB, rL, rE, rO = (f"{TC}!${c}$2:${c}${n_tc}" for c in "BLEO")

    # ---- Sheet 1: Component Inventory ----
    ws = wb.active
    ws.title = S_INV
    write_table(ws, INV_H, inventory, [12, 30, 17, 24, 46, 60, 40, 26, 30, 16, 16, 34],
                wrap=("Purpose", "Dependencies", "Notes", "Existing Tests", "Component Name"), dates=("Last Verified",))
    for r in range(2, len(inventory) + 2):
        a = f"$A{r}"
        mx = f"MAX(_xlfn.MAXIFS({rO},{rB},{a},{rL},\"PASS\"),_xlfn.MAXIFS({rO},{rB},{a},{rL},\"FAIL\"))"
        ws[f"J{r}"] = f"=IF({mx}=0,\"\",{mx})"
        cnt = lambda st: f"COUNTIFS({rB},{a},{rL},\"{st}\")"
        active = f"COUNTIFS({rB},{a},{rE},\"<>{RETIRED.strip()}*\")"
        ws[f"K{r}"] = (f"=IF({active}=0,\"NO TEST CASES\",IF({cnt('FAIL')}>0,\"FAIL\",IF({cnt('BLOCKED')}>0,\"BLOCKED\","
                       f"IF({cnt('PASS')}={active},\"PASS\",IF({cnt('PASS')}>0,\"PARTIAL\",\"NOT RUN\")))))")
        ws[f"J{r}"].number_format = "yyyy-mm-dd hh:mm"
    status_cf(ws, f"K2:K{n_inv}", ["PASS", "FAIL", "BLOCKED", "NOT RUN", "PARTIAL", "NO TEST CASES"])
    ws.conditional_formatting.add(f"A2:L{n_inv}", FormulaRule(formula=[f'$I2="{REMOVED_REQ}"'],
                                                             font=Font(name=ARIAL, color="A6A6A6", strike=True)))

    # ---- Sheet 2: Test Cases ----
    ws = wb.create_sheet(S_TC)
    write_table(ws, TC_H, tcs, [19, 12, 28, 22, 50, 34, 40, 46, 46, 17, 26, 11, 40, 46, 16, 40],
                wrap=("Component Name", "Test Scenario", "Preconditions", "Test Steps", "Expected Result", "Actual Result",
                      "Failure Details", "Evidence / Test Log", "Code or Test Reference"), dates=("Last Executed",))
    dropdown(ws, f"L2:L{n_tc}", STATUSES, "PASS / FAIL need evidence. BLOCKED needs a reason in Failure Details.")
    dropdown(ws, f"J2:J{n_tc}", TEST_TYPES, "Kind of test")
    dropdown(ws, f"K2:K{n_tc}", EXEC_METHODS, "How it is executed")
    status_cf(ws, f"L2:L{n_tc}", STATUSES)
    ws.conditional_formatting.add(f"A2:P{n_tc}", FormulaRule(formula=[f'LEFT($E2,9)="{RETIRED.strip()}"'],
                                                            font=Font(name=ARIAL, color="A6A6A6", italic=True)))

    # ---- Sheet 3: Defects ----
    ws = wb.create_sheet(S_DEF)
    write_table(ws, DEF_H, defects, [10, 24, 30, 50, 40, 40, 50, 10, 11, 40, 36, 12, 16],
                wrap=("Related Test Case ID", "Component", "Description", "Steps to Reproduce", "Expected Behavior",
                      "Actual Behavior", "Related File", "Fix Reference"), dates=("Last Updated",))
    dropdown(ws, f"H2:H{n_def}", SEVERITIES, "Severity")
    dropdown(ws, f"I2:I{n_def}", DEFECT_STATUSES, "OPEN → FIXED → CLOSED (REOPENED if a retest fails)")
    dropdown(ws, f"L2:L{n_def}", ["PASS", "FAIL", "NOT RUN"], "Result of the latest retest")
    status_cf(ws, f"I2:I{n_def}", DEFECT_STATUSES)
    status_cf(ws, f"H2:H{n_def}", SEVERITIES)
    status_cf(ws, f"L2:L{n_def}", ["PASS", "FAIL", "NOT RUN"])

    # ---- Sheet 4: Test Summary ----
    ws = wb.create_sheet(S_SUM)
    write_summary(ws, tcs, inventory, bl_rows, meta, n_tc, n_inv, n_def)

    # ---- extra sheets ----
    ws = wb.create_sheet(S_BL)
    write_table(ws, BL_H, bl_rows, [17, 8, 40, 16, 9, 9, 8, 9, 70, 50, 10], wrap=("Summary", "Evidence / Log", "Command"),
                dates=("Executed At",), freeze="A2")
    dropdown(ws, f"E2:E{max(200, len(bl_rows) + 50)}", ["PASS", "FAIL"], "Exit code 0 = PASS")
    status_cf(ws, f"E2:E{max(200, len(bl_rows) + 50)}", ["PASS", "FAIL"])

    ws = wb.create_sheet(S_RUN)
    write_table(ws, RUN_H, run_rows, [17, 16, 16, 12, 10, 16, 26, 9, 9, 10, 10, 9, 9, 10, 10, 50],
                wrap=("Notes", "Scope"), dates=("Started", "Finished"), freeze="A2")

    ws = wb.create_sheet(S_HIST)
    write_table(ws, HIST_H, hist_rows, [16, 17, 20, 12, 12, 18, 70], wrap=("Note",), dates=("Timestamp",), freeze="A2")
    status_cf(ws, f"D2:E{max(5000, len(hist_rows) + 1000)}", STATUSES)

    write_readme(wb.create_sheet(S_README), meta)
    wb.calculation.fullCalcOnLoad = True
    tmp = WB_PATH + ".tmp.xlsx"
    wb.save(tmp)
    try:
        os.replace(tmp, WB_PATH)
    except PermissionError:
        os.remove(tmp)
        sys.exit(f"Cannot write {WB_PATH}: close it in Excel and run again (nothing was lost; results are in test-evidence/runs).")


def write_summary(ws, tcs, inventory, bl_rows, meta, n_tc, n_inv, n_def):
    TC, INV, DEF = f"'{S_TC}'", f"'{S_INV}'", f"'{S_DEF}'"
    col = lambda sheet, c, n: f"{sheet}!${c}$2:${c}${n}"
    tcL, tcE, tcJ, tcK, tcD, tcB = (col(TC, c, n_tc) for c in "LEJKDB")
    not_ret = f'{tcE},"<>{RETIRED.strip()}*"'
    bold = Font(name=ARIAL, bold=True, size=10)
    ws["A1"] = "Component Testing Summary"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14, color="1F3864")
    ws["A2"] = (f"Last sync {meta['ts']:%Y-%m-%d %H:%M} · run {meta['run_id']} · branch {meta['branch']} @ {meta['commit']}"
                f"{' (uncommitted changes)' if meta['dirty'] else ''} · scope: {meta['scope']}")
    ws["A3"] = "All figures below are formulas over the Test Cases, Component Inventory and Defects sheets. See the README sheet for rules."
    for c in ("A2", "A3"):
        ws[c].font = Font(name=ARIAL, size=9, italic=True, color="595959")

    def header(row, labels):
        for i, lab in enumerate(labels):
            c = ws.cell(row=row, column=1 + i, value=lab)
            c.font, c.fill = HEAD_FONT, HEAD_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)

    header(5, ["Metric", "Value", "How it is calculated"])
    metrics = [
        ("Total components", f"=COUNTA({col(INV, 'A', n_inv)})-COUNTIF({col(INV, 'I', n_inv)},\"{REMOVED_REQ}\")",
         "Inventory rows minus removed files"),
        ("Total test cases", f"=COUNTA({col(TC, 'A', n_tc)})-COUNTIF({tcE},\"{RETIRED.strip()}*\")", "Test Cases rows minus retired"),
        ("PASS", f'=COUNTIFS({tcL},"PASS",{not_ret})', "Status = PASS"),
        ("FAIL", f'=COUNTIFS({tcL},"FAIL",{not_ret})', "Status = FAIL"),
        ("BLOCKED", f'=COUNTIFS({tcL},"BLOCKED",{not_ret})', "Status = BLOCKED"),
        ("NOT RUN", f'=COUNTIFS({tcL},"NOT RUN",{not_ret})', "Status = NOT RUN"),
        ("Overall pass percentage", '=IF((B8+B9)=0,"n/a",B8/(B8+B9))', "PASS ÷ (PASS + FAIL) × 100. BLOCKED and NOT RUN are excluded."),
        ("Execution coverage", '=IF(B7=0,0,(B8+B9)/B7)', "(PASS + FAIL) ÷ total test cases"),
        ("Open defects", f'=COUNTIF({col(DEF, "I", n_def)},"OPEN")+COUNTIF({col(DEF, "I", n_def)},"REOPENED")', "Status OPEN or REOPENED"),
        ("Fixed defects awaiting retest", f'=COUNTIFS({col(DEF, "I", n_def)},"FIXED",{col(DEF, "L", n_def)},"<>PASS")',
         "Status FIXED and Retest Result not PASS"),
        ("Components with ≥1 test case", f'=COUNTA({col(INV, "A", n_inv)})-COUNTIF({col(INV, "K", n_inv)},"NO TEST CASES")',
         "Inventory rows whose Overall Status ≠ NO TEST CASES"),
    ]
    for i, (k, f, how) in enumerate(metrics, start=6):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=2, value=f).font = bold
        ws.cell(row=i, column=3, value=how).font = Font(name=ARIAL, size=9, color="595959")
    ws["B12"].number_format = "0.0%"
    ws["B13"].number_format = "0.0%"
    status_cf(ws, "A8:A11", ["PASS", "FAIL", "BLOCKED", "NOT RUN"])

    # build & lint (separate from functional results)
    r = 18
    header(r, ["Build & Lint (latest)", "Result", "Executed At", "Errors", "Warnings", "Summary"])
    for k, check in enumerate(("Lint", "Build"), start=1):
        idx = max([i for i, b in enumerate(bl_rows) if b["Check"] == check], default=None)
        ws.cell(row=r + k, column=1, value=check).font = bold
        if idx is None:
            ws.cell(row=r + k, column=2, value="NOT RUN")
            continue
        src_row = idx + 2
        for cidx, c in enumerate("EDGHI", start=2):
            cell = ws.cell(row=r + k, column=cidx, value=f"='{S_BL}'!{c}{src_row}")
            if c == "D":
                cell.number_format = "yyyy-mm-dd hh:mm"
    status_cf(ws, f"B{r + 1}:B{r + 2}", ["PASS", "FAIL", "NOT RUN"])

    def breakdown(start, title_txt, rng, values, wildcard=False):
        header(start, [title_txt, "Test cases", "PASS", "FAIL", "BLOCKED", "NOT RUN", "Pass %"])
        for i, v in enumerate(values, start=start + 1):
            crit = f'"{v}*"' if wildcard else f'"{v}"'
            ws.cell(row=i, column=1, value=v).font = BODY_FONT
            ws.cell(row=i, column=2, value=f"=COUNTIFS({rng},{crit},{not_ret})")
            for j, st in enumerate(STATUSES, start=3):
                ws.cell(row=i, column=j, value=f'=COUNTIFS({rng},{crit},{tcL},"{st}",{not_ret})')
            c = ws.cell(row=i, column=7, value=f'=IF((C{i}+D{i})=0,"n/a",C{i}/(C{i}+D{i}))')
            c.number_format = "0.0%"
        return start + len(values) + 2

    r = breakdown(22, "By test type", tcJ, TEST_TYPES)
    r = breakdown(r, "By execution method", tcK, EXEC_METHODS)
    areas = sorted({str(t.get("Feature") or "—").split(" / ")[0] for t in tcs})
    r = breakdown(r, "By area (Module / Feature prefix)", tcD, areas, wildcard=True)

    # component-wise status
    header(r, ["Component ID", "Component", "Module / Feature", "Test cases", "PASS", "FAIL", "BLOCKED", "NOT RUN",
               "Pass %", "Overall Status", "Last Verified"])
    comp_ids = []
    for t in tcs:
        if t.get("Component ID") not in comp_ids and not str(t.get("Test Scenario") or "").startswith(RETIRED):
            comp_ids.append(t.get("Component ID"))
    inv_by_id = {c["Component ID"]: c for c in inventory}
    comp_ids.sort()
    first = r + 1
    for i, cid in enumerate(comp_ids, start=first):
        comp = inv_by_id.get(cid, {})
        ws.cell(row=i, column=1, value=cid)
        ws.cell(row=i, column=2, value=comp.get("Component Name", "—"))
        ws.cell(row=i, column=3, value=comp.get("Module / Feature", "—"))
        ws.cell(row=i, column=4, value=f"=COUNTIFS({tcB},$A{i},{not_ret})")
        for j, st in enumerate(STATUSES, start=5):
            ws.cell(row=i, column=j, value=f'=COUNTIFS({tcB},$A{i},{tcL},"{st}",{not_ret})')
        ws.cell(row=i, column=9, value=f'=IF((E{i}+F{i})=0,"n/a",E{i}/(E{i}+F{i}))').number_format = "0.0%"
        ws.cell(row=i, column=10, value=f"=IFERROR(INDEX({col(INV, 'K', n_inv)},MATCH($A{i},{col(INV, 'A', n_inv)},0)),\"—\")")
        ws.cell(row=i, column=11, value=f"=IFERROR(INDEX({col(INV, 'J', n_inv)},MATCH($A{i},{col(INV, 'A', n_inv)},0)),\"\")").number_format = "yyyy-mm-dd hh:mm"
    last = first + len(comp_ids) - 1
    status_cf(ws, f"J{first}:J{last}", ["PASS", "FAIL", "BLOCKED", "NOT RUN", "PARTIAL", "NO TEST CASES"])
    for row in ws.iter_rows(min_row=6, max_row=last):
        for cell in row:
            if cell.font != HEAD_FONT and not cell.font.bold:
                cell.font = BODY_FONT
    for i, w in enumerate([30, 34, 34, 12, 10, 10, 10, 10, 10, 16, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"


def write_readme(ws, meta):
    lines = [
        ("COMPONENT_TESTING_TRACKER.xlsx — how this workbook is maintained", "title"),
        ("", None),
        ("Single source of truth for component-testing records of the Orange Tree LMS frontend (lms-web).", None),
        ("Regenerated by: python scripts/test-tracker/sync_tracker.py  (npm run test:tracker / npm run test:tracker:full)", None),
        (f"Default targets ({TARGET}): web {BASE_URL} · API {API_URL} (override with TRACKER_BASE_URL / TRACKER_API_URL).", None),
        ("", None),
        ("What runs automatically on each sync", "h"),
        ("• Inventory: every source file under src/ (plus scripts/ui-smoke.mjs) — new files get the next CMP id; deleted files are kept and struck through.", None),
        ("• Unit tests: every src/**/__tests__/*.test.js via node. Each [PASS]/[FAIL] line is one test case (TC-UT-*), matched by '<file> :: <name>'.", None),
        ("• Code-level checks (TC-CL-*): static scans of the working tree for the CLAUDE.md business/architecture rules. They prove source structure, not runtime behaviour.", None),
        ("• HTTP/API checks (TC-API-*, with --http): read-only or rejected-before-side-effect requests to the app/API under test (local by default), with request and response recorded.", None),
        ("• Build & Lint (with --lint / --build): recorded on the 'Build & Lint' sheet, separately from functional results.", None),
        ("• Manual/browser results (TC-UI-*): taken from scripts/test-tracker/manual-results.json (append-only), or typed directly into the Test Cases sheet.", None),
        ("", None),
        ("Evidence rules", "h"),
        ("• PASS/FAIL only from an actual execution in this workbook's history. A file existing, code looking right or a green build is never a PASS.", None),
        ("• UI – Browser cases stay NOT RUN until someone exercises the real UI and records what they saw. BLOCKED always states why in Failure Details.", None),
        ("• Pass % = PASS ÷ (PASS + FAIL) × 100. BLOCKED and NOT RUN never count as passed.", None),
        ("• Testing runs locally (frontend :3000, backend :5000) by default. Results recorded earlier against the deployed Vercel build say so in their Evidence/env; unit and code-level results reflect the local tree at the commit in 'Run Log'.", None),
        ("", None),
        ("Which cells you may edit", "h"),
        ("• Component Inventory: Purpose, Testing Required, Notes (clear Purpose/Testing Required to have them regenerated). Other columns are rewritten each sync; J/K are formulas.", None),
        ("• Test Cases: for TC-UI-* rows, Status / Actual Result / Failure Details / Evidence / Last Executed (always fill Last Executed). Definitions of TC-CL/API/UI rows live in scripts/test-tracker/catalog.py; TC-UT rows come from the test files.", None),
        ("• Defects: every column. The sync only adds defects, and moves Status/Retest Result when a linked automated case is re-run (OPEN→FIXED on PASS, FIXED/CLOSED→REOPENED on FAIL). Close a defect yourself.", None),
        ("• Never renumber IDs. Retired cases are kept with a '[RETIRED]' prefix and excluded from all counts.", None),
        ("", None),
        ("Colour legend", "h"),
        ("PASS", "PASS"), ("FAIL", "FAIL"), ("BLOCKED", "BLOCKED"), ("NOT RUN", "NOT RUN"),
        ("PARTIAL — some test cases passed, the rest not run yet", "PARTIAL"),
        ("", None),
        ("Known limitations", "h"),
        ("• There is no npm test runner in this repo: unit tests are standalone node scripts whose [PASS]/[FAIL] output is parsed.", None),
        ("• Browser cases are not automated; they are recorded by a person (or an agent driving a real browser) in manual-results.json.", None),
        ("• Backend (backend/lms-api, separate repository with node --test) is out of scope: its tests may touch the shared database.", None),
        ("• Mutating flows (create/edit/delete, publish, Razorpay) are NOT RUN by policy on the shared deployment; they need a sandbox course/account.", None),
        ("• The workbook must be closed in Excel while the sync runs (the file is replaced atomically).", None),
    ]
    for i, (text, style) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if style == "title":
            c.font = Font(name=ARIAL, bold=True, size=14, color="1F3864")
        elif style == "h":
            c.font = Font(name=ARIAL, bold=True, size=11, color="1F3864")
        elif style in FILLS:
            bg, fg = FILLS[style]
            c.fill, c.font = PatternFill("solid", fgColor=bg), Font(name=ARIAL, bold=True, color=fg)
        else:
            c.font = BODY_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150


# ==========================================================================
# 5. JSON export for the testing dashboard
# ==========================================================================

DASH_DIR = os.path.join(ROOT, "test-evidence", "dashboard-data")


def _iso(v):
    return v.isoformat() if isinstance(v, dt.datetime) else (v or None)


def export_dashboard(inventory, tcs, defects, bl_rows, run_rows):
    """Write the workbook's records as JSON row files (one per sheet) for the dashboard."""
    os.makedirs(DASH_DIR, exist_ok=True)
    files = {
        "test_cases.json": [{
            "id": t["Test Case ID"], "status": t.get("Status"), "test_type": t.get("Test Type"),
            "method": t.get("Execution Method"), "component_id": t.get("Component ID"),
            "component": t.get("Component Name"), "feature": t.get("Feature"),
            "area": str(t.get("Feature") or "—").split(" / ")[0],
            "scenario": str(t.get("Test Scenario") or "").replace(RETIRED, ""),
            "retired": str(t.get("Test Scenario") or "").startswith(RETIRED),
            "expected": clip(t.get("Expected Result"), 400), "actual": clip(t.get("Actual Result"), 400),
            "failure": clip(t.get("Failure Details"), 600), "evidence": clip(t.get("Evidence / Test Log"), 400),
            "last_executed": _iso(t.get("Last Executed")), "reference": clip(t.get("Code or Test Reference"), 200),
        } for t in tcs],
        "defects.json": [{
            "id": d["Defect ID"], "severity": d.get("Severity"), "status": d.get("Status"),
            "description": d.get("Description"), "related_tc": d.get("Related Test Case ID"),
            "component": d.get("Component"), "actual": clip(d.get("Actual Behavior"), 600),
            "expected": d.get("Expected Behavior"), "file": d.get("Related File"),
            "retest": d.get("Retest Result"), "last_updated": _iso(as_dt(d.get("Last Updated")) or d.get("Last Updated")),
        } for d in defects],
        "components.json": [{
            "id": c["Component ID"], "name": c.get("Component Name"), "type": c.get("Component Type"),
            "feature": c.get("Module / Feature"), "area": str(c.get("Module / Feature") or "—").split(" / ")[0],
            "path": c.get("File Path"), "existing_tests": c.get("Existing Tests"),
            "testing_required": c.get("Testing Required"), "removed": c.get("Testing Required") == REMOVED_REQ,
        } for c in inventory],
        "build_lint.json": [{
            "run_id": b.get("Run ID"), "check": b.get("Check"), "result": b.get("Result"),
            "executed_at": _iso(as_dt(b.get("Executed At")) or b.get("Executed At")), "exit_code": b.get("Exit Code"),
            "errors": b.get("Errors"), "warnings": b.get("Warnings"), "summary": clip(b.get("Summary"), 500),
            "commit": b.get("Git Commit"),
        } for b in bl_rows],
        "runs.json": [{
            "run_id": r.get("Run ID"), "started": _iso(as_dt(r.get("Started")) or r.get("Started")),
            "branch": r.get("Git Branch"), "commit": r.get("Git Commit"), "working_tree": r.get("Working Tree"),
            "scope": r.get("Scope"), "unit_pass": r.get("Unit PASS"), "unit_fail": r.get("Unit FAIL"),
            "code_pass": r.get("Code checks PASS"), "code_fail": r.get("Code checks FAIL"),
            "http_pass": r.get("HTTP PASS"), "http_fail": r.get("HTTP FAIL"), "http_blocked": r.get("HTTP BLOCKED"),
        } for r in run_rows],
    }
    for name, rows in files.items():
        with open(os.path.join(DASH_DIR, name), "w", encoding="utf-8") as fh:
            json.dump(rows, fh, ensure_ascii=False, indent=1, default=str)
    return DASH_DIR


# ==========================================================================
# main
# ==========================================================================

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--http", action="store_true", help="run HTTP/API checks against TRACKER_BASE_URL / TRACKER_API_URL (local by default)")
    ap.add_argument("--lint", action="store_true", help="run npm run lint")
    ap.add_argument("--build", action="store_true", help="run npm run build")
    ap.add_argument("--all", action="store_true", help="--http --lint --build")
    ap.add_argument("--skip-unit", action="store_true", help="do not run unit tests")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if args.all:
        args.http = args.lint = args.build = True

    started = now()
    run_id = "R" + started.strftime("%Y%m%d-%H%M%S")
    run_dir = os.path.join(EVIDENCE_ROOT, run_id)
    os.makedirs(run_dir, exist_ok=True)
    commit = git("rev-parse", "--short", "HEAD")
    branch = git("rev-parse", "--abbrev-ref", "HEAD")
    dirty_files = [l for l in git("status", "--porcelain").splitlines() if l.strip() and "test-evidence/" not in l
                   and "COMPONENT_TESTING_TRACKER" not in l]
    scope = ["inventory", "code checks", "manual results"] + ([] if args.skip_unit else ["unit"]) + \
            (["http"] if args.http else []) + (["lint"] if args.lint else []) + (["build"] if args.build else [])
    print(f"Run {run_id} on {branch}@{commit} ({len(dirty_files)} uncommitted changes) · scope: {', '.join(scope)}")

    wb = load_workbook(WB_PATH) if os.path.exists(WB_PATH) else None
    old_inv, old_tc, old_def = read_rows(wb, S_INV), read_rows(wb, S_TC), read_rows(wb, S_DEF)
    old_bl, old_run, old_hist = read_rows(wb, S_BL), read_rows(wb, S_RUN), read_rows(wb, S_HIST)

    tests = test_files()
    inventory = build_inventory(old_inv, tests, f"{started:%Y-%m-%d}")
    print(f"  inventory: {len(inventory)} components")

    unit_results, unit_files = ([], {}) if args.skip_unit else run_unit_tests(tests, run_dir)
    if not args.skip_unit:
        print(f"  unit: {len(tests)} files, {sum(u['status'] == 'PASS' for u in unit_results)} PASS, "
              f"{sum(u['status'] == 'FAIL' for u in unit_results)} FAIL")
    code_results = run_code_checks(run_dir)
    print(f"  code checks: " + ", ".join(f"{c['id']}={c['status']}" for c in code_results))
    http_results = run_http_checks(run_dir) if args.http else []
    if http_results:
        print(f"  http: " + ", ".join(f"{h['id']}={h['status']}" for h in http_results))

    bl_rows = list(old_bl)
    lint_row = build_row = None
    if args.lint:
        lint_row = run_lint(run_dir, run_id, commit)
        bl_rows.append(lint_row)
        print(f"  lint: {lint_row['Result']} — {lint_row['Summary'][:120]}")
    if args.build:
        build_row = run_build(run_dir, run_id, commit)
        bl_rows.append(build_row)
        print(f"  build: {build_row['Result']} — {build_row['Summary'][:120]}")

    manual_log = []
    if os.path.exists(MANUAL_RESULTS):
        manual_log = json.load(open(MANUAL_RESULTS, encoding="utf-8")).get("results", [])

    hist = list(old_hist)
    tcs, defects, applied = merge(old_tc, inventory, unit_results, unit_files, code_results, http_results, manual_log,
                                  run_id, started, hist, tests, old_def, lint_row)

    finished = now()
    cnt = lambda rows, st: sum(1 for r in rows if r["status"] == st)
    run_rows = list(old_run) + [dict(zip(RUN_H, [
        run_id, started, finished, branch, commit, f"{len(dirty_files)} uncommitted change(s)" if dirty_files else "clean",
        ", ".join(scope), cnt(unit_results, "PASS"), cnt(unit_results, "FAIL"), cnt(code_results, "PASS"),
        cnt(code_results, "FAIL"), cnt(http_results, "PASS"), cnt(http_results, "FAIL"), cnt(http_results, "BLOCKED"),
        applied, f"Evidence: test-evidence/runs/{run_id}/"]))]
    meta = dict(ts=finished, run_id=run_id, branch=branch, commit=commit, dirty=bool(dirty_files), scope=", ".join(scope))
    write_workbook(inventory, tcs, defects, bl_rows, run_rows, hist, meta)
    print(f"  dashboard data: {rel(export_dashboard(inventory, tcs, defects, bl_rows, run_rows))}/")

    active = [t for t in tcs if not str(t.get("Test Scenario") or "").startswith(RETIRED)]
    tally = {s: sum(1 for t in active if t.get("Status") == s) for s in STATUSES}
    rate = tally["PASS"] / (tally["PASS"] + tally["FAIL"]) * 100 if (tally["PASS"] + tally["FAIL"]) else 0
    print(f"Saved {rel(WB_PATH)}: {len(inventory)} components, {len(active)} test cases "
          f"({', '.join(f'{k} {v}' for k, v in tally.items())}; pass {rate:.1f}%), "
          f"{sum(1 for d in defects if d.get('Status') in ('OPEN', 'REOPENED'))} open defects")
    failed_automated = [t for t in active if t.get("Last Executed") == started and t.get("Status") == "FAIL"
                        and str(t["Test Case ID"]).startswith(("TC-UT-",))]
    return 1 if failed_automated else 0


if __name__ == "__main__":
    sys.exit(main())
